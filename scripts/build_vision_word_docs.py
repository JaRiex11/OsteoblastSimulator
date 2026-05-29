# -*- coding: utf-8 -*-
"""Сборка документов Word для папки «Видение проекта» из markdown в docs/."""
from __future__ import annotations

import re
import sys
import shutil
import urllib.parse
import urllib.request
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT = ROOT / "Видение проекта"


def _set_doc_defaults(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    for level, size in [(1, 16), (2, 15), (3, 14)]:
        h = doc.styles[f"Heading {level}"]
        h.font.name = "Times New Roman"
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = RGBColor(0, 0, 0)
        h._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def _add_runs(paragraph, text: str, *, monospace: bool = False) -> None:
    """Разбор **жирного** и `кода` в одной строке."""
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos : m.start()])
        chunk = m.group(0)
        if chunk.startswith("**"):
            run = paragraph.add_run(chunk[2:-2])
            run.bold = True
        else:
            run = paragraph.add_run(chunk[1:-1])
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt(11)
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])
    if monospace and not paragraph.runs:
        run = paragraph.add_run(text)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")


def _set_cell_shading(cell, fill: str = "E8E8E8") -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def _add_table(doc: Document, rows: list[list[str]], header: bool = True) -> None:
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = table.rows[i].cells[j]
            text = row[j] if j < len(row) else ""
            cell.text = ""
            p = cell.paragraphs[0]
            _add_runs(p, text.strip())
            if header and i == 0:
                for run in p.runs:
                    run.bold = True
                _set_cell_shading(cell)
    doc.add_paragraph()


def _parse_table_block(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        if not line.strip().startswith("|"):
            break
        parts = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.match(r"^:?-+:?$", p.replace(" ", "")) for p in parts):
            continue
        rows.append(parts)
    return rows


def md_to_docx(md_path: Path, docx_path: Path, *, title: str | None = None) -> None:
    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = Document()
    _set_doc_defaults(doc)

    if title:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(16)
        doc.add_paragraph()

    i = 0
    in_code = False
    code_lang = ""
    code_lines: list[str] = []

    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("```"):
            lang = line.strip()[3:].strip().lower()
            if not in_code:
                in_code = True
                code_lang = lang
                code_lines = []
                i += 1
                continue
            if code_lang in ("mermaid", "plantuml"):
                code_lines = []
                in_code = False
                i += 1
                continue
            if code_lines:
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(10)
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            code_lines = []
            in_code = False
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if line.strip() == "---":
            i += 1
            continue

        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
            i += 1
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
            i += 1
            continue

        if line.strip().startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            rows = _parse_table_block(block)
            if rows:
                _add_table(doc, rows, header=True)
            continue

        if re.match(r"^\d+\.\s", line.strip()):
            p = doc.add_paragraph(style="List Number")
            _add_runs(p, re.sub(r"^\d+\.\s*", "", line.strip()))
            i += 1
            continue

        if line.strip().startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            _add_runs(p, line.strip()[2:])
            i += 1
            continue

        if line.strip().startswith("*") and line.strip().endswith("*") and not line.strip().startswith("**"):
            p = doc.add_paragraph()
            run = p.add_run(line.strip().strip("*"))
            run.italic = True
            i += 1
            continue

        stripped = line.strip()
        if stripped:
            p = doc.add_paragraph()
            _add_runs(p, stripped)
        i += 1

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)


def _plantuml_encode(text: str) -> str:
    import zlib

    def _encode6bit(b: int) -> str:
        if b < 10:
            return chr(48 + b)
        b -= 10
        if b < 26:
            return chr(65 + b)
        b -= 26
        if b < 26:
            return chr(97 + b)
        b -= 26
        if b == 0:
            return "-"
        if b == 1:
            return "_"
        raise ValueError(b)

    def _append3(b1: int, b2: int, b3: int) -> str:
        c1 = b1 >> 2
        c2 = ((b1 & 0x3) << 4) | (b2 >> 4)
        c3 = ((b2 & 0xF) << 2) | (b3 >> 6)
        c4 = b3 & 0x3F
        return "".join(_encode6bit(c) for c in (c1, c2, c3, c4))

    data = zlib.compress(text.encode("utf-8"))[2:-4]
    out = []
    for i in range(0, len(data), 3):
        chunk = data[i : i + 3]
        if len(chunk) == 3:
            out.append(_append3(chunk[0], chunk[1], chunk[2]))
        elif len(chunk) == 2:
            out.append(_append3(chunk[0], chunk[1], 0))
        else:
            out.append(_append3(chunk[0], 0, 0))
    return "".join(out)


def fetch_plantuml_png(puml_path: Path, png_path: Path) -> bool:
    """Рендер диаграммы через публичный сервер PlantUML."""
    src = puml_path.read_text(encoding="utf-8")
    plantuml = _plantuml_encode(src)
    url = f"https://www.plantuml.com/plantuml/png/{plantuml}"
    try:
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=60) as resp:
            data = resp.read()
        if len(data) < 500 or b"error" in data[:200].lower():
            return False
        png_path.write_bytes(data)
        return True
    except Exception:
        return False


def add_diagram_to_classes_doc(docx_path: Path, png_path: Path | None) -> None:
    if not png_path or not png_path.exists():
        return
    doc = Document(docx_path)
    for para in list(doc.paragraphs):
        t = para.text
        if "Файл PlantUML" in t or "plantuml.com" in t:
            para.clear()
            para.add_run(
                "Диаграмма классов приведена на рисунке 1. Исходник: файл planning_classes.puml "
                "в этой папке."
            )

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(png_path), width=Cm(16))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run("Рисунок 1 — Упрощённая диаграмма классов (планирование)").italic = True
    doc.save(docx_path)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)

    mapping = [
        (DOCS / "Видение_проекта.md", OUT / "1. Видение проекта.docx", "Видение проекта"),
        (DOCS / "Варианты_использования.md", OUT / "2. Варианты использования.docx", None),
        (DOCS / "Классы_предметной_области.md", OUT / "3. Классы предметной области.docx", None),
        (DOCS / "ПОДГОТОВКА_К_РАЗРАБОТКЕ.md", OUT / "0. Содержание комплекта.docx", None),
    ]

    for src, dst, title in mapping:
        if src.exists():
            md_to_docx(src, dst, title=title)
            print(f"OK: {dst.name}")

    puml_src = DOCS / "planning_classes.puml"
    puml_dst = OUT / "planning_classes.puml"
    if puml_src.exists():
        shutil.copy2(puml_src, puml_dst)

    png = OUT / "diagram_planning_classes.png"
    render_script = ROOT / "scripts" / "render_planning_diagram.py"
    if render_script.exists():
        import subprocess
        subprocess.run([sys.executable, str(render_script)], check=False)
    elif puml_src.exists():
        try:
            sys.path.insert(0, str(ROOT / "scripts"))
            from plantuml_render import render_diagram
            render_diagram(puml_src, png, "png")
        except Exception as e:
            print(f"PNG planning: {e}")
    if png.exists():
        add_diagram_to_classes_doc(OUT / "3. Классы предметной области.docx", png)
    else:
        print("PNG: пропущен; в Word — только таблицы классов")

    readme = OUT / "README.txt"
    readme.write_text(
        "Комплект «Подготовка к разработке проекта» (Word).\n\n"
        "0. Содержание комплекта.docx\n"
        "1. Видение проекта.docx\n"
        "2. Варианты использования.docx\n"
        "3. Классы предметной области.docx\n"
        "planning_classes.puml — исходник диаграммы\n"
        "diagram_planning_classes.png — рисунок для вставки в отчёт (если сгенерирован)\n\n"
        "Исходные markdown: папка docs/\n",
        encoding="utf-8",
    )
    print(f"Папка: {OUT}")


if __name__ == "__main__":
    main()
